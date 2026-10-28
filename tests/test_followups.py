"""
Tests for Excel email campaigns (routers/email_campaigns.py, modules/email_campaign_worker.py),
follow-ups (routers/followups.py, modules/followups_core.py, modules/followup_ai.py) and the
company mailbox resolver (modules/company_mailer.py).

    venv\\Scripts\\python -m unittest discover -s tests -t . -v

Ported from serphawk_crm_seo (DaPros uses the same roles and follow_ups columns). Self-contained like
tests/test_inbox.py: in-memory SQLite, a fake SMTP sender, no OpenAI key (so every AI path exercises its
rules / template fallback). tests/__init__.py has already disabled .env loading and dropped the mail / AI
variables, and every engine reference is repointed at SQLite for each test — nothing here can reach
production or send a real email.
"""
import csv
import io
import os
import smtplib
import unittest

from tests._fixtures import brand_as_dapros
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from unittest import mock

from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlmodel import Session, select
from starlette.middleware.base import BaseHTTPMiddleware

from database import (
    ActivityLog, ClientNote, ClientProfile, Contact, ConversationLog, EmailAgentProfile, EmailCampaign,
    EmailCampaignRecipient, EmailSettings, FollowUp, FollowUpNoteScan, InboxMessage, Lead, Notification, SentEmail,
    Tenant, User, UserEmailIntegration,
)
from modules import company_mailer, email_campaign_worker as worker, followup_ai, followups_core as core
from modules.api_tracker import current_salesperson_id
from modules.email_crypto import encrypt_app_password
from routers import email_campaigns as campaigns_router, followups as followups_router
from tests._support import ENGINE, reset_db, use_engine

KEY = "dGVzdC1rZXktZm9yLXVuaXQtdGVzdHMtMzJieXRlcyE="  # test-only Fernet key
NOW = datetime.now(timezone.utc)
ADMIN, SALES, OTHER_SALES, INTERN, ADMIN_T2 = 1, 2, 3, 4, 5


class _Identity(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        uid = request.headers.get("X-User-ID", "")
        current_salesperson_id.set(int(uid) if uid.isdigit() else None)
        return await call_next(request)


def make_app():
    app = FastAPI()
    app.add_middleware(_Identity)
    app.include_router(followups_router.router)
    app.include_router(campaigns_router.router)
    return app


class FakeSender:
    """Stands in for company_mailer.send_via; optionally fails for some recipients."""

    def __init__(self, fail_for=(), error=None):
        self.calls = []
        self.fail_for = set(fail_for)
        self.error = error

    def __call__(self, account, *, to, subject, body, tracking_id=None, message_id=None, extra_headers=None):
        if to in self.fail_for or self.error:
            raise self.error or smtplib.SMTPRecipientsRefused({to: (550, b"no such user")})
        self.calls.append({"from": account.email, "source": account.source, "to": to, "subject": subject,
                           "body": body, "tracking_id": tracking_id})


class DBTestCase(unittest.TestCase):
    def setUp(self):
        p = mock.patch.dict(os.environ, {"EMAIL_CREDENTIAL_ENCRYPTION_KEY": KEY})
        p.start()
        self.addCleanup(p.stop)
        use_engine(self)
        reset_db()

    def seed(self, integration=True, tenant_mailbox=False):
        with Session(ENGINE) as s:
            s.add_all([Tenant(id=1, name="Agency", business_name="DaPros"), Tenant(id=2, name="Other")])
            brand_as_dapros(s)
            s.add_all([
                User(id=ADMIN, tenant_id=1, email="admin@agency.test", password="x", name="Ada Admin", role="Admin"),
                User(id=SALES, tenant_id=1, email="sam@agency.test", password="x", name="Sam Sales", role="SalesManager"),
                User(id=OTHER_SALES, tenant_id=1, email="olga@agency.test", password="x", name="Olga", role="SalesManager"),
                User(id=INTERN, tenant_id=1, email="ian@agency.test", password="x", name="Ian", role="Intern"),
                User(id=ADMIN_T2, tenant_id=2, email="admin2@other.test", password="x", name="Other Admin", role="Admin"),
                User(id=20, tenant_id=1, email="xavier@xco.com", password="x", name="Xavier Xu", role="Client"),
            ])
            s.add(Lead(id=1, tenant_id=1, company_name="Acme", email="jane@acme.com", owner_id=SALES, status="New"))
            s.add(Contact(id=1, tenant_id=1, first_name="Jane", last_name="Doe", full_name="Jane Doe",
                          email="jane@acme.com", lead_id=1))
            s.add(Lead(id=2, tenant_id=1, company_name="Beta", email="bob@beta.com", owner_id=OTHER_SALES))
            s.add(Lead(id=3, tenant_id=2, company_name="Gamma", email="gil@gamma.com"))
            s.add(ClientProfile(id=1, tenant_id=1, companyName="XCo", userId=20, assignedEmployeeId=SALES,
                                contact_person="Xavier Xu"))
            if integration:
                for uid, addr in ((ADMIN, "ada@agency-mail.test"), (SALES, "sam@agency-mail.test")):
                    s.add(UserEmailIntegration(
                        tenant_id=1, user_id=uid, email=addr, encrypted_app_password=encrypt_app_password("app-pw"),
                        provider="gmail", smtp_host="smtp.gmail.com", smtp_port=587, smtp_security="starttls",
                        imap_host="imap.gmail.com"))
                s.add(EmailAgentProfile(tenant_id=1, user_id=ADMIN, agent_name="Ada from DaPros",
                                        signature="Best Regards,\nAda\nDaPros", auto_append_signature=True))
            if tenant_mailbox:
                s.add(EmailSettings(tenant_id=1, smtp_host="mail.agency.test", smtp_port=465, smtp_user="crm@agency.test",
                                    smtp_pass="tenant-pw", from_name="Agency CRM", from_email="crm@agency.test"))
            s.commit()


class ApiTestCase(DBTestCase):
    def setUp(self):
        super().setUp()
        self.seed()
        self.client = TestClient(make_app())
        self.sender = FakeSender()
        p = mock.patch.object(company_mailer, "send_via", self.sender)
        p.start()
        self.addCleanup(p.stop)

    def req(self, method, url, user=ADMIN, **kw):
        return self.client.request(method, url, headers={"X-User-ID": str(user)}, **kw)


# ─────────────────────────────────────────────────────────────────────────────
# AI fallback: dates from notes
# ─────────────────────────────────────────────────────────────────────────────
class HeuristicTests(unittest.TestCase):
    BASE = date(2026, 10, 9)  # a Friday

    def check(self, text, expected):
        r = followup_ai.heuristic_extract(text, self.BASE)
        self.assertEqual(r["due_date"], expected, text)
        self.assertEqual(r["should_follow_up"], expected is not None, text)

    def test_promises_to_respond_land_at_the_end_of_the_period(self):
        self.check("Spoke today. They will discuss internally and respond next week.", "2026-10-16")
        self.check("Please follow up in 2 weeks about pricing", "2026-10-23")
        self.check("Follow up with them in a couple of weeks", "2026-10-23")
        self.check("Will get back next friday", "2026-10-16")

    def test_specific_days_and_dates(self):
        self.check("Client said call back tomorrow", "2026-10-12")  # Saturday rolls to Monday
        self.check("They will review the proposal and let us know on Monday", "2026-10-12")
        self.check("Need budget approval, will confirm by 20 Oct", "2026-10-20")
        self.check("Asked us to check in early next week", "2026-10-12")

    def test_no_commitment_or_no_date_means_no_follow_up(self):
        self.check("Not interested, closed the deal", None)
        self.check("They will discuss internally.", None)  # commitment but no timing — rules can't date it
        self.check("Met them next week at the expo", None)  # date but nothing pending

    def test_ai_output_is_validated(self):
        cleaned = followup_ai._clean_suggestion(
            {"should_follow_up": True, "due_date": "2026-10-10", "channel": "fax", "priority": "urgent",
             "confidence": "9"}, self.BASE, floor=self.BASE, engine="ai")
        self.assertEqual(cleaned["due_date"], "2026-10-12")  # Saturday -> Monday
        self.assertEqual(cleaned["channel"], "email")
        self.assertEqual(cleaned["priority"], "medium")
        self.assertEqual(cleaned["confidence"], 1.0)
        past = followup_ai._clean_suggestion({"should_follow_up": True, "due_date": "2020-01-01"},
                                             self.BASE, floor=self.BASE, engine="ai")
        self.assertEqual(past["due_date"], "2026-10-09")  # never before the floor

    def test_ai_extraction_is_used_and_falls_back_to_rules(self):
        from modules import email_ai
        ai = {"should_follow_up": True, "due_date": "2026-10-16", "title": "Chase decision", "reason": "Discussing internally",
              "quote": "respond next week", "channel": "call", "priority": "high", "confidence": 0.9}
        with mock.patch.object(email_ai, "_chat_json", return_value=ai) as chat:
            r = followup_ai.extract_from_note("They will discuss internally and respond next week.", self.BASE, self.BASE, "Acme")
        self.assertEqual((r["engine"], r["due_date"], r["channel"], r["priority"]), ("ai", "2026-10-16", "call", "high"))
        self.assertIn("<crm_context>", chat.call_args.args[1])  # note fenced as untrusted data
        with mock.patch.object(email_ai, "_chat_json", side_effect=email_ai.EmailAIError("down")):
            r = followup_ai.extract_from_note("They will discuss internally and respond next week.", self.BASE, self.BASE)
        self.assertEqual((r["engine"], r["due_date"]), ("rules", "2026-10-16"))

    def test_ai_personalization_result_is_validated(self):
        from modules import email_ai
        with mock.patch.object(email_ai, "_chat_json", return_value={"subject": "", "body": "Hi Jane, tailored."}):
            p = worker.personalize("Original subject", "Hi Jane", {"Notes": "likes ads"})
        self.assertEqual(p, {"subject": "Original subject", "body": "Hi Jane, tailored."})
        with mock.patch.object(email_ai, "_chat_json", return_value={"subject": "x", "body": " "}):
            with self.assertRaises(email_ai.EmailAIError):
                worker.personalize("s", "b", {})

    def test_template_email_is_personalized_without_ai(self):
        t = followup_ai.template_email(contact_name="Jane Doe", company="Acme", title="Follow up",
                                       reason="pricing", last_discussion="you said you'd discuss internally",
                                       sender_name="Sam")
        self.assertTrue(t["body"].startswith("Hi Jane,"))
        self.assertIn("discuss internally", t["body"])
        self.assertIn("Acme", t["subject"])


# ─────────────────────────────────────────────────────────────────────────────
# Upload parsing + templating
# ─────────────────────────────────────────────────────────────────────────────
def xlsx_bytes(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ParseAndTemplateTests(unittest.TestCase):
    def test_xlsx_with_detected_columns_and_numeric_cells(self):
        data = xlsx_bytes([["First Name", "Last Name", "Email Address", "Company", "Phone"],
                           ["Jane", "Doe", "jane@acme.com", "Acme", 9876543210],
                           [None, None, None, None, None],
                           ["Bob", "", "bob@beta", "Beta", None]])
        parsed = worker.parse_upload("contacts.xlsx", data)
        self.assertEqual(parsed["columns"], ["First Name", "Last Name", "Email Address", "Company", "Phone"])
        self.assertEqual(len(parsed["rows"]), 2)  # blank row dropped
        self.assertEqual(parsed["rows"][0]["Phone"], "9876543210")  # not "9876543210.0"
        self.assertEqual(parsed["detected"]["email"], "Email Address")
        self.assertEqual(parsed["detected"]["first_name"], "First Name")
        self.assertEqual(parsed["detected"]["company"], "Company")
        self.assertEqual(parsed["stats"], {"rows": 2, "sendable": 1, "missing_email": 0, "invalid_email": 1, "duplicates": 0})
        self.assertIn("phone", parsed["placeholders"])

    def test_csv_with_bom_semicolons_and_email_detected_by_content(self):
        text = "﻿Contact;Org;Address\nJane Doe;Acme;jane@acme.com\nJohn;Acme;JANE@acme.com\n"
        parsed = worker.parse_upload("list.csv", text.encode("utf-8"))
        self.assertEqual(parsed["columns"], ["Contact", "Org", "Address"])
        self.assertEqual(parsed["detected"]["email"], "Address")
        self.assertEqual(parsed["detected"]["name"], "Contact")
        self.assertEqual(parsed["stats"]["duplicates"], 1)

    def test_rejects_bad_files(self):
        with self.assertRaises(worker.UploadError):
            worker.parse_upload("old.xls", b"\xd0\xcf\x11\xe0 legacy")
        with self.assertRaises(worker.UploadError):
            worker.parse_upload("x.csv", b"")
        with self.assertRaises(worker.UploadError):
            worker.parse_upload("x.pdf", b"%PDF")
        with mock.patch.dict(os.environ, {"BULK_EMAIL_MAX_ROWS": "2"}):
            with self.assertRaises(worker.UploadError):
                worker.parse_upload("x.csv", b"email\na@x.com\nb@x.com\nc@x.com\n")

    def test_render_placeholders_fallbacks_and_unknowns(self):
        mapping = {"email": "Email", "name": "Full Name", "company": "Company"}
        ctx = worker.row_context({"Full Name": "Jane Doe", "Email": "j@a.com", "Company": "", "City": "Pune"},
                                 mapping, "Ada")
        text, missing = worker.render("Hi {{ first_name }}, {{company|your team}} in {{City}} — {{sender_name}}", ctx)
        self.assertEqual(text, "Hi Jane, your team in Pune — Ada")
        self.assertEqual(missing, [])
        text, missing = worker.render("Hi {{first_name}}, about {{company}}.", worker.row_context({"Email": "x@y.com"}, mapping))
        self.assertEqual(text, "Hi there, about.")
        self.assertEqual(missing, ["company"])
        self.assertEqual(worker.unknown_placeholders(["{{city}} {{nope}}", "{{first_name}}"], ["City"]), ["nope"])


# ─────────────────────────────────────────────────────────────────────────────
# Company mailbox resolution
# ─────────────────────────────────────────────────────────────────────────────
class MailAccountTests(DBTestCase):
    def test_user_integration_wins_then_tenant_mailbox(self):
        self.seed(integration=True, tenant_mailbox=True)
        with Session(ENGINE) as s:
            acct = company_mailer.resolve_account(s, ADMIN, 1)
            self.assertEqual((acct.email, acct.password, acct.source), ("ada@agency-mail.test", "app-pw", "user_integration"))
            self.assertEqual(acct.from_name, "Ada from DaPros")
            acct = company_mailer.resolve_account(s, OTHER_SALES, 1)  # no integration of their own
            self.assertEqual((acct.email, acct.source, acct.smtp_port), ("crm@agency.test", "tenant_settings", 465))
            self.assertNotIn("app-pw", repr(company_mailer.resolve_account(s, ADMIN, 1)))

    def test_never_falls_back_to_the_server_mailbox_for_contact_email(self):
        self.seed(integration=False)
        env = {"EMAIL_SENDER": "fixed@server.test", "EMAIL_PASSWORD": "pw", "SMTP_SERVER": "smtp.server.test"}
        with Session(ENGINE) as s, mock.patch.dict(os.environ, env):
            with self.assertRaises(company_mailer.MailAccountError):
                company_mailer.resolve_account(s, ADMIN, 1)
            status = company_mailer.account_status(s, ADMIN, 1)
            self.assertFalse(status["configured"])
            self.assertIn("Email Integration", status["message"])
            # Internal reminders may use it.
            self.assertEqual(company_mailer.resolve_account(s, ADMIN, 1, allow_server_fallback=True).source, "server_env")

    def test_inactive_integration_is_ignored(self):
        self.seed(integration=True, tenant_mailbox=True)
        with Session(ENGINE) as s:
            row = s.exec(select(UserEmailIntegration).where(UserEmailIntegration.user_id == ADMIN)).one()
            row.is_active = False
            s.add(row)
            s.commit()
            self.assertEqual(company_mailer.resolve_account(s, ADMIN, 1).source, "tenant_settings")

    def test_signature_applied_once(self):
        sig = "Best Regards,\nAda\nDaPros"
        body = company_mailer.apply_signature("Hi Jane,\n\nThanks!\n\nBest regards,\nAda\nada@dapros.com.mx", sig)
        self.assertEqual(body.count("Best Regards"), 1)
        self.assertTrue(body.endswith(sig))
        self.assertEqual(company_mailer.apply_signature(body, sig), body)

    def test_plain_closing_is_replaced_when_the_signature_has_its_own(self):
        sig = "Best Regards,\nAda Lovelace\nEmail: ada@agency.test"
        for body in ("Hi Jane,\n\nShall we talk?\n\nBest regards,\nAda Lovelace",
                     "Hi Jane,\n\nShall we talk?\n\nKind regards,\n\nAda\nE2E Digital"):
            out = company_mailer.apply_signature(body, sig)
            self.assertEqual(out, "Hi Jane,\n\nShall we talk?\n\n" + sig)
        # A signature without a closing line keeps the body's own closing.
        out = company_mailer.apply_signature("Hi,\n\nBest regards,\nAda", "Ada Lovelace | +91 900")
        self.assertIn("Best regards,\nAda\n\nAda Lovelace | +91 900", out)
        # Body text that merely starts with "Thanks" is not a closing.
        self.assertIn("Thanks for your time today.", company_mailer.apply_signature("Hi,\n\nThanks for your time today.", sig))


# ─────────────────────────────────────────────────────────────────────────────
# Campaign API + worker
# ─────────────────────────────────────────────────────────────────────────────
ROWS = [
    {"First Name": "Jane", "Email": "jane@acme.com", "Company": "Acme", "City": "Pune"},
    {"First Name": "", "Email": "new@delta.io", "Company": "Delta", "City": "Delhi"},
    {"First Name": "Dup", "Email": "JANE@acme.com", "Company": "Acme", "City": ""},
    {"First Name": "Bad", "Email": "not-an-email", "Company": "X", "City": ""},
    {"First Name": "Empty", "Email": "", "Company": "Y", "City": ""},
]
COLUMNS = ["First Name", "Email", "Company", "City"]
MAPPING = {"email": "Email", "first_name": "First Name", "company": "Company"}


def campaign_payload(**over):
    body = {"columns": COLUMNS, "rows": ROWS, "mapping": MAPPING, "subject": "Growth ideas for {{company}}",
            "body": "Hi {{first_name}},\n\nWe help businesses in {{city|your city}} grow.\n\nBest regards,\nAda",
            "name": "October outreach", "source_filename": "october.xlsx", "create_leads": True, "followup_days": 3}
    body.update(over)
    return body


class CampaignApiTests(ApiTestCase):
    def create(self, user=ADMIN, **over):
        with mock.patch.object(worker, "start_campaign") as start:
            res = self.req("POST", "/email-campaigns", user=user, json=campaign_payload(**over))
        return res, start

    def run_worker(self, cid, sender=None):
        worker.run_campaign(cid, sender=sender or self.sender, delay=0)

    def test_parse_endpoint_accepts_xlsx_upload(self):
        data = xlsx_bytes([["Name", "Email"], ["Jane Doe", "jane@acme.com"]])
        res = self.req("POST", "/email-campaigns/parse", files={"file": ("list.xlsx", data, "application/octet-stream")})
        self.assertEqual(res.status_code, 200, res.text)
        self.assertEqual(res.json()["detected"]["email"], "Email")
        bad = self.req("POST", "/email-campaigns/parse", files={"file": ("list.xls", b"\xd0\xcf", "application/octet-stream")})
        self.assertEqual(bad.status_code, 400)
        self.assertIn(".xlsx", bad.json()["detail"])

    def test_sender_status_names_the_company_account(self):
        res = self.req("GET", "/email-campaigns/sender").json()
        self.assertEqual((res["configured"], res["address"], res["source"]), (True, "ada@agency-mail.test", "user_integration"))
        self.assertNotIn("password", str(res).lower())

    def test_preview_renders_each_row_and_lists_skips(self):
        res = self.req("POST", "/email-campaigns/preview", json=campaign_payload(count=5))
        self.assertEqual(res.status_code, 200, res.text)
        data = res.json()
        self.assertEqual(data["sendable"], 2)
        first, second = data["previews"]
        self.assertEqual(first["subject"], "Growth ideas for Acme")
        self.assertTrue(first["body"].startswith("Hi Jane,"))
        self.assertIn("in Pune grow", first["body"])
        self.assertTrue(second["body"].startswith("Hi there,"))
        self.assertTrue(first["body"].endswith("Best Regards,\nAda\nDaPros"))  # Email Agent signature
        self.assertEqual(sorted(s["reason"].split()[0] for s in data["skipped"]), ["Duplicate", "No", "“not-an-email”"])

    def test_unknown_placeholder_and_missing_email_column_are_rejected(self):
        res = self.req("POST", "/email-campaigns/preview", json=campaign_payload(subject="Hi {{nickname}}"))
        self.assertEqual(res.status_code, 400)
        self.assertIn("{{nickname}}", res.json()["detail"])
        res = self.req("POST", "/email-campaigns/preview", json=campaign_payload(mapping={"email": None}))
        self.assertEqual(res.status_code, 400)

    def test_campaign_sends_personalized_emails_from_the_company_account_and_records_everything(self):
        res, start = self.create()
        self.assertEqual(res.status_code, 200, res.text)
        cid = res.json()["campaign"]["id"]
        start.assert_called_once_with(cid)
        self.assertEqual(res.json()["campaign"]["status"], "Queued")
        self.run_worker(cid)

        self.assertEqual([c["to"] for c in self.sender.calls], ["jane@acme.com", "new@delta.io"])
        self.assertTrue(all(c["from"] == "ada@agency-mail.test" for c in self.sender.calls))
        self.assertEqual(self.sender.calls[0]["subject"], "Growth ideas for Acme")
        self.assertTrue(self.sender.calls[0]["body"].startswith("Hi Jane,"))
        self.assertTrue(self.sender.calls[1]["body"].startswith("Hi there,"))
        self.assertEqual(self.sender.calls[0]["body"].lower().count("best regards"), 1)  # template closing replaced by signature

        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual((c.status, c.total, c.sent_count, c.failed_count, c.skipped_count), ("Completed", 5, 2, 0, 3))
            recips = {r.row_number: r for r in s.exec(select(EmailCampaignRecipient)).all()}
            self.assertEqual([recips[i].status for i in range(1, 6)], ["Sent", "Sent", "Skipped", "Skipped", "Skipped"])
            self.assertIsNotNone(recips[1].sent_at)
            self.assertEqual(recips[1].lead_id, 1)  # linked to the existing Acme lead, not duplicated
            sent = s.exec(select(SentEmail).order_by(SentEmail.id)).all()
            self.assertEqual([(e.to_email, e.status, e.tenant_id) for e in sent],
                             [("jane@acme.com", "Sent", 1), ("new@delta.io", "Sent", 1)])
            self.assertEqual(self.sender.calls[0]["tracking_id"], sent[0].id)  # open tracking pixel id
            new_lead = s.get(Lead, recips[2].lead_id)
            self.assertEqual((new_lead.email, new_lead.source, new_lead.company_name, new_lead.owner_id),
                             ("new@delta.io", "Email Campaign", "Delta", ADMIN))
            self.assertEqual(s.get(Lead, 1).status, "Contacted")
            fus = s.exec(select(FollowUp).order_by(FollowUp.id)).all()
            self.assertEqual(len(fus), 2)
            self.assertEqual({f.source for f in fus}, {"campaign"})
            self.assertEqual([f.assigned_to for f in fus], [SALES, ADMIN])  # Jane's lead owner; the new lead's owner (sender)
            self.assertEqual(core.as_utc(fus[0].due_at).date(), (NOW + timedelta(days=3)).date())
            self.assertTrue(s.exec(select(Notification).where(Notification.user_id == ADMIN)).first())
            self.assertTrue(s.exec(select(ActivityLog).where(ActivityLog.lead_id == 1)).first())

        detail = self.req("GET", f"/email-campaigns/{cid}?status=Sent").json()
        self.assertEqual([r["email"] for r in detail["recipients"]], ["jane@acme.com", "new@delta.io"])
        export = self.req("GET", f"/email-campaigns/{cid}/export")
        self.assertEqual(export.status_code, 200)
        lines = list(csv.reader(io.StringIO(export.text.lstrip("﻿"))))
        self.assertEqual(lines[0][:5], ["Row", "Recipient email", "Name", "Company", "Status"])
        self.assertEqual([l[4] for l in lines[1:]], ["Sent", "Sent", "Skipped", "Skipped", "Skipped"])

    def test_requires_a_configured_company_account(self):
        with mock.patch.dict(os.environ, {"EMAIL_SENDER": "fixed@server.test", "EMAIL_PASSWORD": "pw",
                                          "SMTP_SERVER": "smtp.server.test"}):
            res, start = self.create(user=OTHER_SALES)  # no integration, tenant has no mailbox
        self.assertEqual(res.status_code, 503)
        self.assertIn("Email Integration", res.json()["detail"])
        start.assert_not_called()

    def test_failed_rows_are_recorded_and_can_be_retried(self):
        res, _ = self.create(followup_days=None)
        cid = res.json()["campaign"]["id"]
        self.run_worker(cid, FakeSender(fail_for={"new@delta.io"}))
        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual((c.status, c.sent_count, c.failed_count), ("Completed", 1, 1))
            failed = s.exec(select(EmailCampaignRecipient).where(EmailCampaignRecipient.status == "Failed")).one()
            self.assertIn("refused this recipient", failed.error)
            self.assertEqual(len(s.exec(select(SentEmail)).all()), 1)  # failed send leaves no sent_emails row
            self.assertIsNone(failed.lead_id)  # ...and no lead created just for the bounced address
            self.assertIsNone(s.exec(select(Lead).where(Lead.email == "new@delta.io")).first())
        with mock.patch.object(worker, "start_campaign") as start:
            retry = self.req("POST", f"/email-campaigns/{cid}/retry")
        self.assertEqual(retry.status_code, 200, retry.text)
        self.assertEqual(retry.json()["requeued"], 1)
        start.assert_called_once_with(cid)
        self.run_worker(cid)
        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual((c.status, c.sent_count, c.failed_count), ("Completed", 2, 0))

    def test_login_failure_stops_the_whole_campaign(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        self.run_worker(cid, FakeSender(error=smtplib.SMTPAuthenticationError(535, b"bad credentials")))
        with Session(ENGINE) as s:
            c = s.get(EmailCampaign, cid)
            self.assertEqual(c.status, "Failed")
            self.assertIn("app password", c.last_error)
            statuses = sorted(r.status for r in s.exec(select(EmailCampaignRecipient)).all())
            self.assertEqual(statuses.count("Failed"), 1)
            self.assertEqual(statuses.count("Pending"), 1)  # never attempted

    def test_cancel_stops_before_the_next_row(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        self.assertEqual(self.req("POST", f"/email-campaigns/{cid}/cancel").status_code, 200)
        self.run_worker(cid)
        self.assertEqual(self.sender.calls, [])
        self.assertEqual(self.req("GET", f"/email-campaigns/{cid}").json()["campaign"]["pending"], 2)

    def test_interrupted_rows_are_failed_not_resent(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        with Session(ENGINE) as s:
            r = s.exec(select(EmailCampaignRecipient).where(EmailCampaignRecipient.row_number == 1)).one()
            r.status, r.claimed_at = "Sending", NOW - timedelta(minutes=30)
            s.add(r)
            s.commit()
        with mock.patch.object(worker, "start_campaign", return_value=True) as start:
            self.assertEqual(worker.resume_campaigns(), 1)
        start.assert_called_once()
        self.assertEqual(start.call_args.args[0], cid)
        self.run_worker(cid)
        self.assertEqual([c["to"] for c in self.sender.calls], ["new@delta.io"])
        with Session(ENGINE) as s:
            r = s.exec(select(EmailCampaignRecipient).where(EmailCampaignRecipient.row_number == 1)).one()
            self.assertEqual(r.status, "Failed")
            self.assertIn("Interrupted", r.error)

    def test_campaigns_are_scoped_by_role_and_tenant(self):
        res, _ = self.create()
        cid = res.json()["campaign"]["id"]
        self.assertEqual(self.req("GET", f"/email-campaigns/{cid}", user=SALES).status_code, 404)
        self.assertEqual(self.req("GET", "/email-campaigns", user=SALES).json()["items"], [])
        self.assertEqual(self.req("GET", f"/email-campaigns/{cid}", user=ADMIN_T2).status_code, 404)
        self.assertEqual(self.req("GET", "/email-campaigns", user=INTERN).status_code, 403)
        self.assertEqual(self.req("GET", "/email-campaigns", user=None).status_code, 401)
        self.assertEqual(len(self.req("GET", "/email-campaigns").json()["items"]), 1)

    def test_demo_trial_limit_counts_every_email(self):
        with Session(ENGINE) as s:
            s.add(User(id=9, tenant_id=1, email="demo@agency.test", password="x", name="Demo", role="Demo"))
            s.add(UserEmailIntegration(tenant_id=1, user_id=9, email="demo@agency-mail.test",
                                       encrypted_app_password=encrypt_app_password("pw"), smtp_host="smtp.x", imap_host="imap.x"))
            t = s.get(Tenant, 1)
            t.is_trial, t.limit_emails, t.usage_emails = True, 5, 4
            s.add(t)
            s.commit()
        res, start = self.create(user=9)
        self.assertEqual(res.status_code, 403)
        self.assertEqual(res.json()["detail"]["error"], "LIMIT_REACHED")
        start.assert_not_called()


# ─────────────────────────────────────────────────────────────────────────────
# Follow-up API
# ─────────────────────────────────────────────────────────────────────────────
def iso(dt):
    return dt.isoformat()


class FollowUpApiTests(ApiTestCase):
    def create(self, user=SALES, **body):
        payload = {"lead_id": 1, "title": "Call about pricing", "due_at": iso(NOW + timedelta(days=2)), "tz_offset": 0}
        payload.update(body)
        return self.req("POST", "/followups", user=user, json=payload)

    def test_create_on_a_lead_fills_contact_and_owner(self):
        res = self.create(notes="They will discuss internally")
        self.assertEqual(res.status_code, 200, res.text)
        item = res.json()["item"]
        self.assertEqual((item["contact_name"], item["contact_email"], item["company_name"]), ("Jane Doe", "jane@acme.com", "Acme"))
        self.assertEqual(item["assigned_to"]["id"], SALES)
        self.assertEqual(item["state"], "upcoming")
        self.assertEqual(item["lead"], {"id": 1, "name": "Acme"})
        with Session(ENGINE) as s:
            self.assertEqual(s.exec(select(ActivityLog).where(ActivityLog.lead_id == 1)).one().action, "Follow-up scheduled")

    def test_date_only_due_uses_the_viewers_timezone(self):
        res = self.create(due_at="2026-11-02", tz_offset=-330)  # IST viewer
        self.assertEqual(res.json()["item"]["due_at"], "2026-11-02T04:30:00+00:00")  # 10:00 IST

    def test_access_rules(self):
        self.assertEqual(self.create(user=OTHER_SALES).status_code, 403)  # not their lead
        self.assertEqual(self.create(user=INTERN).status_code, 403)
        self.assertEqual(self.create(user=ADMIN_T2).status_code, 404)  # other tenant
        self.assertEqual(self.create(lead_id=None, contact_email="").status_code, 400)
        fid = self.create().json()["item"]["id"]
        self.assertEqual(self.req("GET", "/followups?view=all", user=OTHER_SALES).json()["total"], 0)
        self.assertEqual(self.req("GET", "/followups?view=all", user=ADMIN).json()["total"], 1)
        self.assertEqual(self.req("PATCH", f"/followups/{fid}", user=OTHER_SALES, json={"title": "x"}).status_code, 404)
        self.assertEqual(self.req("GET", "/followups?view=all", user=ADMIN_T2).json()["total"], 0)

    def test_views_and_summary_separate_overdue_today_upcoming_completed(self):
        start_today = datetime(NOW.year, NOW.month, NOW.day, tzinfo=timezone.utc)
        ids = {}
        for key, due in (("overdue", start_today - timedelta(days=2)), ("today", start_today + timedelta(hours=12)),
                         ("upcoming", start_today + timedelta(days=4)), ("done", start_today - timedelta(days=1))):
            ids[key] = self.create(title=key, due_at=iso(due)).json()["item"]["id"]
        self.assertEqual(self.req("POST", f"/followups/{ids['done']}/complete", user=SALES,
                                  json={"outcome": "Spoke, they signed"}).status_code, 200)
        for view, expected in (("overdue", ["overdue"]), ("today", ["today"]), ("upcoming", ["upcoming"]),
                               ("completed", ["done"]), ("pending", ["overdue", "today", "upcoming"])):
            got = [i["title"] for i in self.req("GET", f"/followups?view={view}&tz_offset=0", user=SALES).json()["items"]]
            self.assertEqual(got, expected, view)
        summary = self.req("GET", "/followups/summary?tz_offset=0", user=SALES).json()
        self.assertEqual({k: summary["counts"][k] for k in ("pending", "overdue", "today", "upcoming", "completed")},
                         {"pending": 3, "overdue": 1, "today": 1, "upcoming": 1, "completed": 1})
        self.assertEqual([i["title"] for i in summary["next"]], ["overdue", "today", "upcoming"])
        states = {i["title"]: i["state"] for i in self.req("GET", "/followups?view=all&tz_offset=0", user=SALES).json()["items"]}
        self.assertEqual(states, {"overdue": "overdue", "today": "today", "upcoming": "upcoming", "done": "completed"})

    def test_reschedule_resets_reminder_and_cancel_reopen(self):
        fid = self.create().json()["item"]["id"]
        with Session(ENGINE) as s:
            fu = s.get(FollowUp, fid)
            fu.reminder_sent_at = NOW
            s.add(fu)
            s.commit()
        res = self.req("PATCH", f"/followups/{fid}", user=SALES, json={"due_at": iso(NOW + timedelta(days=9)), "tz_offset": 0})
        self.assertEqual(res.status_code, 200, res.text)
        item = res.json()["item"]
        self.assertIsNone(item["reminder_sent_at"])
        self.assertEqual(item["remind_at"], item["due_at"])
        self.assertEqual(self.req("PATCH", f"/followups/{fid}", user=SALES, json={"status": "Cancelled"}).json()["item"]["state"], "cancelled")
        self.assertEqual(self.req("PATCH", f"/followups/{fid}", user=SALES, json={"status": "Pending"}).json()["item"]["status"], "Pending")
        self.assertEqual(self.req("PATCH", f"/followups/{fid}", user=SALES, json={"auto_send": True}).status_code, 400)

    def test_complete_with_next_follow_up_keeps_the_history_chain(self):
        fid = self.create().json()["item"]["id"]
        res = self.req("POST", f"/followups/{fid}/complete", user=SALES,
                       json={"outcome": "Asked to call again after budget meeting", "next": {"due_at": "2026-12-01"}, "tz_offset": 0})
        self.assertEqual(res.status_code, 200, res.text)
        data = res.json()
        self.assertEqual(data["item"]["status"], "Completed")
        self.assertEqual(data["next"]["source"], "completed_chain")
        self.assertIn("budget meeting", data["next"]["notes"])
        self.assertEqual(self.req("POST", f"/followups/{fid}/complete", user=SALES, json={}).status_code, 400)
        history = self.req("GET", "/followups/history?lead_id=1", user=SALES).json()
        self.assertEqual(len(history["followups"]), 2)
        self.assertIn("followup_done", [e["type"] for e in history["timeline"]])

    def test_draft_from_previous_notes_and_send_from_company_account(self):
        with Session(ENGINE) as s:
            s.add(ClientNote(tenant_id=1, lead_id=1, content="Jane said they will discuss internally and respond next week.",
                             author_name="Sam", created_at=NOW - timedelta(days=1)))
            s.commit()
        fid = self.create(notes="Pricing discussion pending").json()["item"]["id"]
        draft = self.req("POST", f"/followups/{fid}/draft-email", user=SALES, json={})
        self.assertEqual(draft.status_code, 200, draft.text)
        d = draft.json()
        self.assertEqual(d["engine"], "template")  # no OPENAI_API_KEY in tests
        self.assertTrue(d["body"].startswith("Hi Jane,"))
        self.assertIn("discuss internally", d["body"])
        sent = self.req("POST", f"/followups/{fid}/send-email", user=SALES,
                        json={"subject": d["subject"], "body": d["body"], "tz_offset": 0})
        self.assertEqual(sent.status_code, 200, sent.text)
        self.assertEqual(sent.json()["from"], "sam@agency-mail.test")
        self.assertEqual(sent.json()["item"]["status"], "Completed")
        self.assertEqual(self.sender.calls[0]["to"], "jane@acme.com")
        with Session(ENGINE) as s:
            row = s.exec(select(SentEmail)).one()
            self.assertEqual((row.lead_id, row.status, row.tenant_id), (1, "Sent", 1))
            self.assertEqual(s.exec(select(ConversationLog).where(ConversationLog.lead_id == 1)).one().type, "email")
            self.assertEqual(s.get(Lead, 1).status, "Contacted")

    def test_send_without_company_account_is_refused(self):
        with Session(ENGINE) as s:
            s.add(Lead(id=7, tenant_id=1, company_name="Olga's", email="o@x.com", owner_id=OTHER_SALES))
            s.commit()
        fid = self.create(user=OTHER_SALES, lead_id=7).json()["item"]["id"]
        res = self.req("POST", f"/followups/{fid}/send-email", user=OTHER_SALES, json={"subject": "Hi", "body": "Hello"})
        self.assertEqual(res.status_code, 503)
        self.assertEqual(self.sender.calls, [])

    def test_suggest_uses_the_latest_commitment_in_the_notes(self):
        with Session(ENGINE) as s:
            s.add(ConversationLog(tenant_id=1, client_id=1, title="Call with Xavier", type="call",
                                  description="Xavier will check with his partner and get back in 2 weeks.",
                                  created_at=NOW))
            s.commit()
        res = self.req("POST", "/followups/suggest", user=SALES, json={"client_id": 1, "tz_offset": 0})
        self.assertEqual(res.status_code, 200, res.text)
        s_ = res.json()["suggestion"]
        expected = followup_ai.parse_relative_date("get back in 2 weeks", NOW.date())
        self.assertEqual(s_["due_date"], expected.isoformat())
        self.assertTrue(res.json()["email"]["body"].startswith("Hi Xavier,"))
        self.assertEqual(self.req("POST", "/followups/suggest", user=OTHER_SALES, json={"client_id": 1}).status_code, 403)

    def test_scan_endpoint_creates_follow_up_from_a_note(self):
        with Session(ENGINE) as s:
            s.add(ClientNote(id=5, tenant_id=1, lead_id=1, content="They will discuss internally and respond next week.",
                             created_at=NOW))
            s.commit()
        res = self.req("POST", "/followups/scan", user=SALES, json={"lead_id": 1, "tz_offset": 0})
        self.assertEqual(res.status_code, 200, res.text)
        created = res.json()["created"]
        self.assertEqual(len(created), 1)
        self.assertEqual(created[0]["source"], "ai_auto")
        self.assertEqual(created[0]["source_ref"], "note-5")
        self.assertEqual(self.req("POST", "/followups/scan", user=SALES, json={"lead_id": 1}).json()["scanned"], 0)


# ─────────────────────────────────────────────────────────────────────────────
# Scheduler: reminders, auto-send, AI note scan
# ─────────────────────────────────────────────────────────────────────────────
class SchedulerTests(DBTestCase):
    def setUp(self):
        super().setUp()
        self.seed(integration=True)

    def add_fu(self, **kw):
        base = dict(tenant_id=1, lead_id=1, title="Call Jane", due_at=NOW - timedelta(minutes=5), assigned_to=SALES,
                    created_by=SALES, contact_name="Jane Doe", contact_email="jane@acme.com", company_name="Acme")
        base.update(kw)
        with Session(ENGINE) as s:
            fu = FollowUp(**base)
            s.add(fu)
            s.commit()
            return fu.id

    def test_due_reminder_is_sent_exactly_once(self):
        due = self.add_fu(notes="They promised an answer this week")
        self.add_fu(title="Later", due_at=NOW + timedelta(days=1))
        self.add_fu(title="Done", status="Completed")
        sender = FakeSender()
        self.assertEqual(core.process_due_reminders(send_email=sender), 1)
        self.assertEqual(core.process_due_reminders(send_email=sender), 0)
        self.assertEqual(len(sender.calls), 1)
        self.assertEqual(sender.calls[0]["to"], "sam@agency.test")
        self.assertEqual(sender.calls[0]["from"], "sam@agency-mail.test")
        self.assertIn("Call Jane", sender.calls[0]["subject"])
        self.assertIn("They promised an answer", sender.calls[0]["body"])
        with Session(ENGINE) as s:
            notes = s.exec(select(Notification)).all()
            self.assertEqual([(n.user_id, n.link) for n in notes], [(SALES, "/leads/1?tab=followups")])
            self.assertIsNotNone(s.get(FollowUp, due).reminder_sent_at)

    def test_custom_remind_time_fires_before_due(self):
        self.add_fu(due_at=NOW + timedelta(days=1), remind_at=NOW - timedelta(minutes=1))
        self.assertEqual(core.process_due_reminders(send_email=FakeSender()), 1)

    def test_auto_send_sends_prepared_email_and_completes(self):
        fid = self.add_fu(auto_send=True, email_subject="Checking in", email_body="Hi Jane,\n\nAny update?")
        self.add_fu(title="Not yet", auto_send=True, email_subject="x", email_body="y", due_at=NOW + timedelta(hours=3))
        sender = FakeSender()
        self.assertEqual(core.process_auto_sends(sender=sender), 1)
        self.assertEqual(core.process_auto_sends(sender=sender), 0)
        self.assertEqual([(c["to"], c["from"]) for c in sender.calls], [("jane@acme.com", "sam@agency-mail.test")])
        with Session(ENGINE) as s:
            fu = s.get(FollowUp, fid)
            self.assertEqual(fu.status, "Completed")
            self.assertIsNotNone(fu.sent_email_id)
            self.assertEqual(s.get(SentEmail, fu.sent_email_id).to_email, "jane@acme.com")

    def test_auto_send_failure_is_reported_and_not_retried(self):
        fid = self.add_fu(auto_send=True, email_subject="Checking in", email_body="Hi")
        self.assertEqual(core.process_auto_sends(sender=FakeSender(fail_for={"jane@acme.com"})), 0)
        with Session(ENGINE) as s:
            fu = s.get(FollowUp, fid)
            self.assertEqual((fu.status, fu.auto_send, fu.email_sent_at), ("Pending", False, None))
            self.assertIn("refused", fu.last_error)
            self.assertEqual(s.exec(select(Notification)).one().title, "Follow-up email was not sent")
            self.assertEqual(s.exec(select(SentEmail)).all(), [])

    def test_note_scan_creates_dated_follow_up_once(self):
        with Session(ENGINE) as s:
            s.add(ClientNote(id=1, tenant_id=1, lead_id=1, author_id=OTHER_SALES,
                             content="Call went well. They will discuss internally and respond next week.", created_at=NOW))
            s.add(ClientNote(id=2, tenant_id=1, lead_id=1, content="Sent them our brochure.", created_at=NOW))
            s.commit()
        result = core.scan_notes()
        self.assertEqual(result["scanned"], 2)
        self.assertEqual(len(result["created"]), 1)
        self.assertEqual(core.scan_notes()["scanned"], 0)  # each note is read once
        with Session(ENGINE) as s:
            fu = s.exec(select(FollowUp)).one()
            expected = followup_ai.parse_relative_date("respond next week", NOW.date())
            # The fixture is a DaPros workspace: its 10am is in Mexico City.
            promised = datetime(expected.year, expected.month, expected.day, 10,
                                tzinfo=ZoneInfo("America/Mexico_City")).astimezone(timezone.utc)
            self.assertEqual(core.as_utc(fu.due_at), promised)
            self.assertEqual((fu.source, fu.source_ref, fu.assigned_to, fu.lead_id), ("ai_auto", "note-1", SALES, 1))
            self.assertEqual(fu.contact_email, "jane@acme.com")
            self.assertIn("respond next week", fu.ai_reason)
            self.assertEqual(s.exec(select(Notification)).one().user_id, SALES)
            self.assertEqual(len(s.exec(select(FollowUpNoteScan)).all()), 2)

    def test_note_scan_skips_duplicates_own_emails_spam_and_stale_commitments(self):
        due = core.at_local_hour(followup_ai.parse_relative_date("respond next week", NOW.date()))
        existing = self.add_fu(due_at=due)
        with Session(ENGINE) as s:
            s.add(ClientNote(tenant_id=1, lead_id=1, content="They will respond next week.", created_at=NOW))
            s.add(SentEmail(tenant_id=1, lead_id=1, to_email="jane@acme.com", subject="Follow up",
                            english_body="Let me know next week", sent_at=NOW))
            s.add(ConversationLog(tenant_id=1, lead_id=1, title="Follow up", type="email",
                                  description="Please let us know next week", created_at=NOW))
            s.add(InboxMessage(tenant_id=1, mailbox="m", message_id="<s>", from_address="spam@x.com", lead_id=1,
                               subject="Win", body_text="Reply tomorrow to confirm your prize", ai_category="spam",
                               received_at=NOW))
            s.add(ClientNote(tenant_id=1, client_id=1, content="Will call back tomorrow",
                             created_at=NOW - timedelta(days=20)))  # outside the scan window
            s.commit()
        result = core.scan_notes()
        self.assertEqual(result["created"], [])
        self.assertEqual(result["matched"], [existing])
        with Session(ENGINE) as s:
            self.assertEqual(len(s.exec(select(FollowUp)).all()), 1)
            fu = s.get(FollowUp, existing)
            self.assertEqual((fu.source, fu.title, core.as_utc(fu.due_at)), ("manual", "Call Jane", due))  # a person's choice stays
            self.assertIn("respond next week", fu.ai_reason)  # ...but gains what was said
            self.assertIn("AI (", fu.notes)

    def test_note_moves_a_matching_campaign_follow_up_to_the_promised_date(self):
        day = followup_ai.parse_relative_date("respond next week", NOW.date())
        promised = datetime(day.year, day.month, day.day, 10,
                            tzinfo=ZoneInfo("America/Mexico_City")).astimezone(timezone.utc)
        campaign_fu = self.add_fu(source="campaign", title="Follow up on “Growth ideas”", due_at=promised - timedelta(days=2),
                                  reminder_sent_at=NOW)
        with Session(ENGINE) as s:
            s.add(ClientNote(tenant_id=1, lead_id=1, content="Jane: we'll discuss internally and respond next week.", created_at=NOW))
            s.commit()
        self.assertEqual(core.scan_notes()["matched"], [campaign_fu])
        with Session(ENGINE) as s:
            fu = s.get(FollowUp, campaign_fu)
            self.assertEqual((core.as_utc(fu.due_at), fu.source, fu.reminder_sent_at), (promised, "ai_auto", None))
            self.assertTrue(fu.title.startswith("Follow up:"))

    def test_inbound_email_reply_can_create_a_follow_up(self):
        with Session(ENGINE) as s:
            s.add(InboxMessage(tenant_id=1, mailbox="m", message_id="<r>", from_address="xavier@xco.com", from_name="Xavier",
                               client_id=1, subject="Re: proposal",
                               body_text="Thanks — we'll review it with the board and get back to you next week.",
                               received_at=NOW))
            s.commit()
        result = core.scan_notes()
        self.assertEqual(len(result["created"]), 1)
        with Session(ENGINE) as s:
            fu = s.get(FollowUp, result["created"][0])
            self.assertEqual((fu.client_id, fu.source_ref, fu.contact_email), (1, f"inbox-1", "xavier@xco.com"))


if __name__ == "__main__":
    unittest.main()
